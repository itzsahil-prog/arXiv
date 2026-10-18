"""Loading documents into the reference library."""
import hashlib
import io
import json
import os
import time
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zlib

from .db import get_params
from .engine import kgram_hashes, tokenize, winnow

DEFAULT_K, DEFAULT_W = 5, 4
ATOM = "{http://www.w3.org/2005/Atom}"


class NullCtx:
    """Stand-in progress reporter for the command line and tests."""

    def tick(self, processed, message=None, target=None):
        pass

    def cancelled(self):
        return False


# ------------------------------------------------------------------ documents

def read_upload(name, data):
    """Extract plain text from an uploaded .txt, .md, .docx or .pdf."""
    ext = os.path.splitext(name)[1].lower()
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="ignore")
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(data)).pages)
    raise ValueError(f"{name}: unsupported file type. Use .txt, .md, .docx or .pdf.")


def add_doc(con, source, ref, title, url, text, k, w):
    """Store a document and its fingerprints. Returns False if skipped or already present."""
    toks, _ = tokenize(text)
    if len(toks) < k + w:
        return False
    cur = con.execute(
        "INSERT OR IGNORE INTO docs(source, ref, title, url, words, added, text) VALUES (?,?,?,?,?,?,?)",
        (source, ref, title, url, len(toks), time.time(), zlib.compress(text.encode("utf-8"))),
    )
    if cur.rowcount == 0:
        return False
    doc_id = cur.lastrowid
    fps = winnow(kgram_hashes(toks, k), w)
    con.executemany("INSERT INTO fp VALUES (?,?,?)", ((h, doc_id, pos) for pos, h in fps))
    con.execute(
        "INSERT INTO sources(source, docs, words) VALUES (?,1,?) "
        "ON CONFLICT(source) DO UPDATE SET docs=docs+1, words=words+excluded.words",
        (source, len(toks)),
    )
    return True


def delete_doc(con, doc_id):
    row = con.execute("SELECT source, words, text FROM docs WHERE id=?", (doc_id,)).fetchone()
    if not row:
        return False
    k, w = get_params(con)
    toks, _ = tokenize(zlib.decompress(row[2]).decode("utf-8"))
    for _, h in winnow(kgram_hashes(toks, k), w):
        con.execute("DELETE FROM fp WHERE h=? AND doc=?", (h, doc_id))
    con.execute("DELETE FROM docs WHERE id=?", (doc_id,))
    con.execute("UPDATE sources SET docs=docs-1, words=words-? WHERE source=?", (row[1], row[0]))
    con.commit()
    return True


# -------------------------------------------------------------------- sources

def ingest_files(con, files, ctx=None):
    """files: list of (filename, bytes). Returns {added, skipped, failed}."""
    ctx = ctx or NullCtx()
    k, w = get_params(con, DEFAULT_K, DEFAULT_W)
    added = skipped = 0
    failed = []
    for i, (name, data) in enumerate(files, 1):
        if ctx.cancelled():
            break
        try:
            text = read_upload(name, data)
            ref = hashlib.sha1(data).hexdigest()[:16]
            if add_doc(con, "files", ref, name, "", text, k, w):
                added += 1
            else:
                skipped += 1
        except Exception as e:  # one bad file should not stop the batch
            failed.append({"name": name, "error": str(e)})
        ctx.tick(i)
    con.commit()
    return {"added": added, "skipped": skipped, "failed": failed}


def ingest_arxiv_file(con, path, limit=0, ctx=None):
    """Kaggle arXiv metadata JSONL (arxiv-metadata-oai-snapshot.json): title and abstract."""
    ctx = ctx or NullCtx()
    k, w = get_params(con, DEFAULT_K, DEFAULT_W)
    n = seen = 0
    with open(path, encoding="utf-8") as f:
        for line in f:
            if ctx.cancelled():
                break
            seen += 1
            rec = json.loads(line)
            title = " ".join(rec.get("title", "").split())
            abstract = " ".join(rec.get("abstract", "").split())
            if len(abstract.split()) >= 20:
                aid = rec["id"]
                if add_doc(con, "arxiv", aid, title, f"https://arxiv.org/abs/{aid}",
                           f"{title}. {abstract}", k, w):
                    n += 1
            if seen % 500 == 0:
                con.commit()
                ctx.tick(n, f"{n} abstracts added")
            if limit and n >= limit:
                break
    con.commit()
    return f"Added {n} arXiv abstracts."


def parse_arxiv_atom(xml_bytes):
    """Parse an arXiv API Atom response into [{id, title, abstract}]."""
    root = ET.fromstring(xml_bytes)
    out = []
    for entry in root.findall(ATOM + "entry"):
        url = (entry.findtext(ATOM + "id") or "").strip()
        if not url or "/api/errors" in url:
            continue
        aid = url.rsplit("/abs/", 1)[-1]
        if "v" in aid and aid.rsplit("v", 1)[-1].isdigit():
            aid = aid.rsplit("v", 1)[0]
        out.append({
            "id": aid,
            "title": " ".join((entry.findtext(ATOM + "title") or "").split()),
            "abstract": " ".join((entry.findtext(ATOM + "summary") or "").split()),
        })
    return out


def _fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "overlap-similarity-checker/1.0"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return resp.read()


def ingest_arxiv_search(con, query, limit=200, ctx=None, delay=3.2):
    """Pull abstracts from the public arXiv API. It asks for one request every 3 seconds."""
    ctx = ctx or NullCtx()
    k, w = get_params(con, DEFAULT_K, DEFAULT_W)
    query = query.strip()
    if ":" not in query:  # plain words become: all:word AND all:word
        query = " AND ".join("all:" + t for t in query.split())
    n, start, page = 0, 0, 100
    while n < limit and not ctx.cancelled():
        params = urllib.parse.urlencode({
            "search_query": query, "start": start,
            "max_results": min(page, limit - n), "sortBy": "relevance",
        })
        url = "https://export.arxiv.org/api/query?" + params
        for attempt in range(3):
            try:
                records = parse_arxiv_atom(_fetch(url))
                break
            except Exception:
                if attempt == 2:
                    raise
                time.sleep(delay * (attempt + 2))
        if not records:
            break
        for rec in records:
            if len(rec["abstract"].split()) >= 20 and add_doc(
                con, "arxiv", rec["id"], rec["title"], f"https://arxiv.org/abs/{rec['id']}",
                f"{rec['title']}. {rec['abstract']}", k, w,
            ):
                n += 1
        con.commit()
        start += len(records)
        ctx.tick(n, f"{n} abstracts added")
        if len(records) < min(page, limit):
            break
        time.sleep(delay)
    return f"Added {n} arXiv abstracts for \"{query}\"."


def ingest_wikipedia(con, limit=20000, config="20231101.en", ctx=None):
    """Stream Wikipedia articles from the Hugging Face dataset wikimedia/wikipedia."""
    ctx = ctx or NullCtx()
    from datasets import load_dataset
    k, w = get_params(con, DEFAULT_K, DEFAULT_W)
    ds = load_dataset("wikimedia/wikipedia", config, split="train", streaming=True)
    n = seen = 0
    for rec in ds:
        if ctx.cancelled():
            break
        seen += 1
        if add_doc(con, "wikipedia", rec["id"], rec["title"], rec["url"], rec["text"], k, w):
            n += 1
        if seen % 200 == 0:
            con.commit()
            ctx.tick(n, f"{n} articles added")
        if limit and n >= limit:
            break
    con.commit()
    return f"Added {n} Wikipedia articles."


def library_stats(con):
    from .db import INDEX_DB
    try:
        k, w = get_params(con)
    except Exception:
        k = w = None
    sources = [dict(source=r[0], docs=r[1], words=r[2])
               for r in con.execute("SELECT source, docs, words FROM sources WHERE docs > 0 ORDER BY docs DESC")]
    size = sum(os.path.getsize(p) for p in (INDEX_DB, INDEX_DB + "-wal") if os.path.exists(p))
    fingerprints = con.execute("SELECT COALESCE(MAX(rowid), 0) FROM fp").fetchone()[0]
    return {
        "documents": sum(s["docs"] for s in sources),
        "words": sum(s["words"] for s in sources),
        "fingerprints": fingerprints,
        "bytes": size,
        "sources": sources,
        "params": {"k": k, "w": w},
    }
